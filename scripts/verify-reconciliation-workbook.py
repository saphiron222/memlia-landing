# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Recette dans un vrai moteur de tableur, sans simuler les valeurs des formules."""
import json
import shutil
import subprocess
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'public/downloads/modele-rapprochement-bancaire.xlsx'
WORK = ROOT / '.qa/workbook-d'
INPUT = WORK / 'input'
OUTPUT = WORK / 'recalculated'
for folder in (INPUT, OUTPUT):
    folder.mkdir(parents=True, exist_ok=True)

cases = [
    ('concordant', {}, [1100, 1100, 0], 'Soldes concordants'),
    ('ecart', {'B9': 900}, [1100, 1050, 50], 'NON VALIDÉ'),
    ('inexplique', {'B17': 20}, [1100, 1100, 0], 'NON VALIDÉ'),
    ('decouvert', {'B8': -100, 'B9': -250}, [0, -100, 100], 'NON VALIDÉ'),
    ('centimes', {'B8': 0.3, 'B9': 0.1, 'B11': 0, 'B12': 0, 'B14': 0, 'B15': 0.2}, [0.3, 0.3, 0], 'Soldes concordants'),
    ('periode-inversee', {'B5': '2026-02-01'}, None, 'NON VALIDÉ'),
    ('incomplet', {'B8': None}, None, 'NON VALIDÉ'),
    ('sous-centime', {'B8': 1000.001}, None, 'NON VALIDÉ'),
    ('hors-borne', {'B8': 1000000000}, None, 'NON VALIDÉ'),
]
for name, changes, expected, status in cases:
    wb = load_workbook(SOURCE)
    ws = wb['Exemple fictif']
    for cell, value in changes.items():
        if cell == 'B5':
            from datetime import date
            value = date.fromisoformat(value)
        ws[cell] = value
    wb.save(INPUT / f'{name}.xlsx')
shutil.copy2(SOURCE, INPUT / 'modele-rapprochement-bancaire.xlsx')
command = ['soffice', f'-env:UserInstallation={ (WORK / "profile").as_uri() }', '--headless', '--convert-to', 'xlsx', '--outdir', str(OUTPUT), *map(str, INPUT.glob('*.xlsx'))]
result = subprocess.run(command, capture_output=True, text=True, check=True, timeout=180)
print(result.stdout)
results = []
for name, changes, expected, status in cases:
    ws = load_workbook(OUTPUT / f'{name}.xlsx', data_only=True)['Exemple fictif']
    actual = [ws[c].value for c in ('B19', 'B20', 'B22')]
    assert expected is None or actual == expected, (name, actual, expected)
    assert str(ws['B24'].value).startswith(status), (name, ws['B24'].value)
    formula_ws = load_workbook(OUTPUT / f'{name}.xlsx')['Exemple fictif']
    assert formula_ws['B22'].data_type == 'f'
    for cell, value in changes.items():
        if cell != 'B5':
            assert ws[cell].value == value, (name, cell, ws[cell].value, value)
    results.append({'case': name, 'values': actual, 'state': ws['B24'].value, 'pass': True})
blank = load_workbook(OUTPUT / 'modele-rapprochement-bancaire.xlsx', data_only=True)['À remplir']
assert blank['B24'].value.startswith('NON VALIDÉ')
results.append({'case': 'feuille-vide', 'state': blank['B24'].value, 'pass': True})
report = {'engine': subprocess.check_output(['soffice', '--version'], text=True).strip(), 'results': results}
(ROOT / 'docs/qa/site-tools-d').mkdir(parents=True, exist_ok=True)
(ROOT / 'docs/qa/site-tools-d/workbook-recipe.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
print(json.dumps(report, ensure_ascii=False, indent=2))
