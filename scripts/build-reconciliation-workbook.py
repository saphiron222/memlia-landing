# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Construire le modèle statique : uv run scripts/build-reconciliation-workbook.py."""
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'public/downloads/modele-rapprochement-bancaire.xlsx'
INPUTS = (5, 6, 8, 9, 11, 12, 14, 15, 17)
LABELS = {
    5: 'Début de période', 6: 'Fin de période',
    8: 'Solde du relevé bancaire', 9: 'Solde du compte 512',
    11: 'Remises en circulation (+ banque)', 12: 'Paiements en circulation (− banque)',
    14: 'Frais bancaires non comptabilisés (− 512)', 15: 'Intérêts non comptabilisés (+ 512)',
    17: 'Éléments inexpliqués (contrôle séparé)',
    19: 'Solde bancaire ajusté', 20: 'Solde comptable ajusté',
    22: 'Différence banque − comptabilité', 24: 'État du contrôle',
}

def build():
    wb = Workbook()
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True)
    notice = wb.active
    assert notice is not None
    notice.title = 'Notice'
    lines = [
        'Modèle de rapprochement bancaire — contrôle de soldes',
        'Sans macro, sans connexion externe, sans inscription. Travaillez sur une copie locale.',
        'Exemple fictif : jeu illustratif. À remplir : saisies vides, formules conservées.',
        'Saisissez uniquement les cellules jaunes. Les cellules vertes contiennent les formules.',
        'B5/B6 : période du contrôle. B8/B9 : soldes à la même date de fin.',
        'Soldes : positif = avoir disponible / solde débiteur du 512 ; négatif = découvert / solde créditeur du 512.',
        'Remises : déjà au 512, pas encore au relevé. Elles augmentent le solde bancaire ajusté.',
        'Paiements : déjà au 512, pas encore au relevé. Ils diminuent le solde bancaire ajusté.',
        'Frais/intérêts : déjà au relevé, pas encore au 512. Frais soustraits, intérêts ajoutés au 512.',
        'Ajustements : montants positifs par défaut ; une contrepassation porte le signe opposé. Ne pas compter deux fois une même opération.',
        'Éléments inexpliqués : montant à investiguer, non intégré aux soldes. Toute valeur non nulle maintient NON VALIDÉ.',
        'Saisissez 0 si un ajustement est absent ; une cellule vide signifie saisie incomplète.',
        'Formules : banque + remises − paiements ; 512 − frais + intérêts ; différence des deux.',
        'Arrondi à deux décimales. Montants limités à ±999 999 999,99 €, saisie au centime.',
        'Un écart nul indique seulement des soldes concordants : validation humaine toujours nécessaire.',
        'Ce fichier n’apparie aucune ligne, ne justifie aucune opération et ne passe aucune écriture.',
        'Conservez les justificatifs, détaillez les opérations en circulation et les exceptions dans les lignes dédiées.',
        'Référence : ANC, Plan comptable général — https://www.anc.gouv.fr/plan-comptable-general-0',
        'Les formules de ce modèle sont une convention de contrôle de soldes, pas une prescription du PCG.',
    ]
    for row, text in enumerate(lines, 1):
        notice.cell(row, 1, text)
        notice.cell(row, 1).alignment = Alignment(wrap_text=True, vertical='top')
        notice.row_dimensions[row].height = 40
    notice.column_dimensions['A'].width = 100
    notice.sheet_properties.pageSetUpPr.fitToPage = True
    notice.page_setup.orientation = 'portrait'
    notice.page_setup.paperSize = notice.PAPERSIZE_A4
    notice.page_setup.fitToWidth = 1
    notice.page_setup.fitToHeight = 1
    notice.print_area = f'A1:A{len(lines)}'
    for name, example in [('Exemple fictif', True), ('À remplir', False)]:
        ws = wb.create_sheet(name)
        ws['A1'] = 'Contrôle de soldes — ' + name
        ws['A2'] = 'Montants en euros (€). Saisies jaunes ; formules vertes. Voir Notice pour les signes et les limites.'
        ws.merge_cells('A1:D1')
        ws.merge_cells('A2:D2')
        for row, label in LABELS.items():
            ws.cell(row, 1, label)
        if example:
            for row, value in {5: date(2026, 1, 1), 6: date(2026, 1, 31), 8: 1000, 9: 950, 11: 200, 12: 100, 14: 25, 15: 175, 17: 0}.items():
                ws.cell(row, 2, value)
        for row in INPUTS:
            cell = ws.cell(row, 2)
            cell.fill = PatternFill('solid', fgColor='FFF2CC')
            cell.number_format = 'dd/mm/yyyy' if row in (5, 6) else '#,##0.00;[Red]-#,##0.00'
        for row, formula in {19: '=ROUND(B8+B11-B12,2)', 20: '=ROUND(B9-B14+B15,2)', 22: '=ROUND(B19-B20,2)'}.items():
            ws.cell(row, 2, formula)
            ws.cell(row, 2).number_format = '#,##0.00;[Red]-#,##0.00'
        invalid_money = ','.join(f'ABS(B{row})>999999999.99,ROUND(B{row},2)<>B{row}' for row in INPUTS[2:])
        ws['B24'] = f'=IFERROR(IF(COUNT(B5:B6,B8:B9,B11:B12,B14:B15,B17)<>9,"NON VALIDÉ — saisie incomplète",IF(OR({invalid_money}),"NON VALIDÉ — montant hors borne ou précision",IF(B5>B6,"NON VALIDÉ — période inversée",IF(OR(B17<>0,B22<>0),"NON VALIDÉ — écart à expliquer","Soldes concordants — à valider")))),"NON VALIDÉ — saisie invalide")'
        ws.merge_cells('B24:D24')
        for row in (19, 20, 22, 24):
            ws.cell(row, 2).fill = PatternFill('solid', fgColor='E2F0D9')
        ws.conditional_formatting.add('B22', CellIsRule(operator='notEqual', formula=['0'], fill=PatternFill('solid', fgColor='FFC7CE')))
        ws.conditional_formatting.add('B17', CellIsRule(operator='notEqual', formula=['0'], fill=PatternFill('solid', fgColor='FFC7CE')))
        money = DataValidation(type='custom', formula1='AND(ISNUMBER(B8),ABS(B8)<=999999999.99,ROUND(B8,2)=B8)', allow_blank=True)
        money.errorTitle = 'Montant au centime requis'
        money.error = 'Saisissez un nombre au centime dans la borne indiquée dans Notice.'
        money.showErrorMessage = True
        ws.add_data_validation(money)
        for row in INPUTS[2:]:
            money.add(ws.cell(row, 2))
        dates = DataValidation(type='date', operator='between', formula1='DATE(1900,1,1)', formula2='DATE(9999,12,31)', allow_blank=True)
        dates.showErrorMessage = True
        ws.add_data_validation(dates)
        dates.add('B5:B6')
        ws.append([])
        for col, value in enumerate(['Date', 'Libellé / justificatif', 'Type / signe', 'Montant'], 1):
            ws.cell(27, col, value)
        ws['A26'] = 'Détail manuel des opérations et exceptions (non totalisé automatiquement)'
        for row in range(28, 48):
            for col in range(1, 5):
                ws.cell(row, col).fill = PatternFill('solid', fgColor='FFF2CC')
            ws.cell(row, 1).number_format = 'dd/mm/yyyy'
            ws.cell(row, 4).number_format = '#,##0.00;[Red]-#,##0.00'
        ws.column_dimensions['A'].width = 54
        ws.column_dimensions['B'].width = 38
        ws.column_dimensions['C'].width = 24
        ws.column_dimensions['D'].width = 20
        ws.row_dimensions[24].height = 36
        ws.freeze_panes = 'B8'
        ws.print_options.horizontalCentered = True
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.print_area = 'A1:D47'
    for ws in wb:
        for row in ws:
            for cell in row:
                cell.font = Font(name='Calibri', size=11, color='231F20', bold=cell.row == 1)
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.sheet_view.showGridLines = False
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(OUTPUT)

if __name__ == '__main__':
    build()
